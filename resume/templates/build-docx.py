#!/usr/bin/env python3
"""Convert the constrained resume Markdown -> ATS-safe .docx.
Single column, standard headings, List-Bullet style, no tables/text-boxes/images.
Usage: build-docx.py <input.md> <output.docx>"""
import sys, re
from docx import Document
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x55, 0x55, 0x55)
FONT = "Calibri"

inp, outp = sys.argv[1], sys.argv[2]
md = open(inp, encoding="utf-8").read()
m = re.search(r"^#\s+\S", md, re.M)
if m:
    md = md[m.start():]
lines = md.split("\n")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)        # A4
sec.top_margin = sec.bottom_margin = Mm(12)
sec.left_margin = sec.right_margin = Mm(13)
CONTENT_W = Inches((210 - 26) / 25.4)                      # usable width for right tab

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.05

def set_bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), "C9C9C9")
    pbdr.append(bottom); pPr.append(pbdr)

def add_runs(p, text, size=10, bold=False, italic=False, color=None):
    text = text.replace("`", "")
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        r = p.add_run(part[2:-2] if (part.startswith("**") and part.endswith("**")) else part)
        r.bold = bold or (part.startswith("**") and part.endswith("**"))
        r.italic = italic
        r.font.size = Pt(size); r.font.name = FONT
        if color is not None:
            r.font.color.rgb = color

i, cur = 0, ""
while i < len(lines):
    line = lines[i].strip()
    if line == "":
        i += 1; continue

    if re.match(r"^#\s+\S", line) and not line.startswith("##"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line[2:].strip().upper())
        r.bold = True; r.font.size = Pt(19); r.font.name = FONT; r.font.color.rgb = NAVY
        i += 1
        while i < len(lines) and lines[i].strip() == "":
            i += 1
        if i < len(lines) and re.match(r"^\*\*.+\*\*$", lines[i].strip()):
            rp = doc.add_paragraph(); rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rp.paragraph_format.space_after = Pt(1)
            add_runs(rp, lines[i].strip(), size=11, bold=True, color=NAVY)
            i += 1
            while i < len(lines) and lines[i].strip() == "":
                i += 1
        if i < len(lines) and not lines[i].lstrip().startswith("#"):
            cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(4)
            add_runs(cp, lines[i].strip(), size=9, color=GRAY)
            i += 1
        continue

    if line.startswith("## "):
        cur = line[3:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(3)
        r = p.add_run(cur.upper())
        r.bold = True; r.font.size = Pt(11); r.font.name = FONT; r.font.color.rgb = NAVY
        set_bottom_border(p)
        i += 1; continue

    if line.startswith("### "):
        title = line[4:].strip()
        dates = ""
        j = i + 1
        while j < len(lines) and lines[j].strip() == "":
            j += 1
        if j < len(lines) and re.match(r"^\*.+\*$", lines[j].strip()):
            dates = lines[j].strip()[1:-1].strip()
            i = j + 1
        else:
            i = i + 1
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.tab_stops.add_tab_stop(CONTENT_W, WD_TAB_ALIGNMENT.RIGHT)
        add_runs(p, title, size=10.5, bold=True)
        if dates:
            r = p.add_run("\t" + dates)
            r.italic = True; r.font.size = Pt(9); r.font.name = FONT; r.font.color.rgb = GRAY
        continue

    if re.match(r"^[-*]\s+", line):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.line_spacing = 1.05
        add_runs(p, re.sub(r"^[-*]\s+", "", line), size=10)
        i += 1; continue

    if re.match(r"^\*.+\*$", line):
        p = doc.add_paragraph()
        add_runs(p, line[1:-1], size=9, italic=True, color=GRAY)
        i += 1; continue

    # plain paragraph (Summary)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    add_runs(p, line, size=10)
    i += 1

doc.save(outp)
print("DOCX: " + outp)
